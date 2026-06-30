from __future__ import annotations

import argparse
import json
import logging
import re
import sys
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Any

from docx import Document
from openpyxl import Workbook
from tqdm import tqdm


BACKEND_DIR = Path(__file__).resolve().parents[1]
PROJECT_ROOT = BACKEND_DIR.parent
if str(BACKEND_DIR) not in sys.path:
    sys.path.insert(0, str(BACKEND_DIR))

from app.services.gemini import generate_gemini_text, validate_gemini_config  # noqa: E402


ALLOWED_CATEGORIES = {
    "Social",
    "Education",
    "Family",
    "Friendship",
    "Success",
    "Failure",
    "Leadership",
    "Honesty",
    "Wisdom",
    "Patience",
    "Money",
    "Time",
    "Work",
    "Kindness",
    "Responsibility",
    "Morality",
    "General",
}
DEFAULT_BATCH_SIZE = 10
MAX_RETRIES = 3
LOG_FILE = BACKEND_DIR / "metadata_generation.log"


@dataclass(frozen=True)
class ProverbMeaningPair:
    """A source proverb and the meaning at the same paragraph index."""

    row_number: int
    proverb: str
    meaning: str


@dataclass(frozen=True)
class MetadataRecord:
    """Generated metadata for one proverb row."""

    category: str
    keywords: list[str]
    english_translation: str


def configure_logging(log_path: Path) -> None:
    """Configure file logging for metadata generation errors and retries."""

    logging.basicConfig(
        filename=log_path,
        level=logging.INFO,
        format="%(asctime)s %(levelname)s %(message)s",
        encoding="utf-8",
    )


def read_docx_paragraphs(path: Path) -> list[str]:
    """Read non-empty paragraphs from a DOCX file."""

    if not path.exists():
        raise FileNotFoundError(
            f"Required DOCX file not found: {path}. "
            "Place the file there or pass a custom path with --proverbs/--meanings."
        )

    document = Document(path)
    return [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]


def merge_dataset(proverbs: list[str], meanings: list[str]) -> list[ProverbMeaningPair]:
    """Merge proverb and meaning paragraphs by matching index."""

    if len(proverbs) != len(meanings):
        raise ValueError("Number of proverbs does not match number of meanings.")

    return [
        ProverbMeaningPair(row_number=index, proverb=proverb, meaning=meaning)
        for index, (proverb, meaning) in enumerate(zip(proverbs, meanings), start=1)
    ]


def chunked(items: list[ProverbMeaningPair], batch_size: int) -> list[list[ProverbMeaningPair]]:
    """Split source rows into fixed-size batches."""

    return [items[index : index + batch_size] for index in range(0, len(items), batch_size)]


def build_prompt(batch: list[ProverbMeaningPair]) -> str:
    """Build a batch prompt that asks Gemini to return a JSON array."""

    rows = [
        {
            "row_number": item.row_number,
            "proverb": item.proverb,
            "meaning": item.meaning,
        }
        for item in batch
    ]
    rows_json = json.dumps(rows, ensure_ascii=False, indent=2)
    categories = ", ".join(sorted(ALLOWED_CATEGORIES))

    return f"""
You are an expert in Myanmar language, linguistics, and Myanmar proverbs.

Given each proverb and its meaning, generate:

- one category
- between five and ten semantic keywords
- a natural English translation

Allowed categories:
{categories}

Keywords must include semantic concepts instead of only copying words from the proverb.

Return ONLY valid JSON.
Return a JSON array with one object per input row, in the same order.
Each object must have exactly these fields:
- row_number
- category
- keywords
- english_translation

Example object:
{{
  "row_number": 1,
  "category": "Social",
  "keywords": [
    "ချီးမွမ်း",
    "ကဲ့ရဲ့",
    "ဝေဖန်",
    "အပြစ်တင်",
    "ဂုဏ်သတင်း",
    "လူမှုဆက်ဆံရေး"
  ],
  "english_translation": "Praise lasts seven days; criticism lasts seven days."
}}

Input rows:
{rows_json}
""".strip()


def extract_json_array(text: str) -> list[Any]:
    """Parse a JSON array from Gemini text, including fenced-code responses."""

    cleaned = re.sub(r"```(?:json)?|```", "", text.strip(), flags=re.IGNORECASE)
    start = cleaned.find("[")
    end = cleaned.rfind("]")
    if start == -1 or end == -1 or end <= start:
        raise ValueError("Gemini response did not contain a JSON array.")
    parsed = json.loads(cleaned[start : end + 1])
    if not isinstance(parsed, list):
        raise ValueError("Gemini response JSON was not an array.")
    return parsed


def sanitize_metadata(raw: Any, expected_row_number: int) -> MetadataRecord:
    """Validate and normalize one generated metadata object."""

    if not isinstance(raw, dict):
        raise ValueError(f"Row {expected_row_number}: metadata item is not an object.")

    category = str(raw.get("category") or "General").strip()
    if category not in ALLOWED_CATEGORIES:
        logging.warning(
            "Row %s returned unsupported category %r; using General.",
            expected_row_number,
            category,
        )
        category = "General"

    keywords_value = raw.get("keywords")
    if isinstance(keywords_value, list):
        keywords = [str(keyword).strip() for keyword in keywords_value if str(keyword).strip()]
    elif isinstance(keywords_value, str):
        keywords = [keyword.strip() for keyword in keywords_value.split(",") if keyword.strip()]
    else:
        keywords = []

    if len(keywords) < 5 or len(keywords) > 10:
        logging.warning(
            "Row %s returned %s keywords; expected 5-10.",
            expected_row_number,
            len(keywords),
        )

    english_translation = str(raw.get("english_translation") or "").strip()

    return MetadataRecord(
        category=category,
        keywords=keywords[:10],
        english_translation=english_translation,
    )


def fallback_metadata() -> MetadataRecord:
    """Return metadata defaults used after repeated Gemini failure."""

    return MetadataRecord(category="General", keywords=[], english_translation="")


def generate_batch_metadata(batch: list[ProverbMeaningPair]) -> list[MetadataRecord]:
    """Generate metadata for one batch, retrying failed Gemini calls up to three times."""

    prompt = build_prompt(batch)
    last_error: Exception | None = None

    for attempt in range(1, MAX_RETRIES + 1):
        try:
            raw_text = generate_gemini_text(prompt)
            parsed = extract_json_array(raw_text)
            if len(parsed) != len(batch):
                raise ValueError(
                    f"Gemini returned {len(parsed)} records for {len(batch)} input rows."
                )

            records = [
                sanitize_metadata(raw_item, source.row_number)
                for raw_item, source in zip(parsed, batch)
            ]
            return records
        except (json.JSONDecodeError, ValueError) as exc:
            last_error = exc
            logging.exception(
                "Batch rows %s-%s JSON/parsing error on attempt %s/%s.",
                batch[0].row_number,
                batch[-1].row_number,
                attempt,
                MAX_RETRIES,
            )
        except Exception as exc:
            last_error = exc
            logging.exception(
                "Batch rows %s-%s API error on attempt %s/%s.",
                batch[0].row_number,
                batch[-1].row_number,
                attempt,
                MAX_RETRIES,
            )

        if attempt < MAX_RETRIES:
            time.sleep(2**attempt)

    logging.error(
        "Failed rows %s-%s after %s attempts: %s",
        batch[0].row_number,
        batch[-1].row_number,
        MAX_RETRIES,
        last_error,
    )
    return [fallback_metadata() for _ in batch]


def save_workbook(
    output_path: Path,
    pairs: list[ProverbMeaningPair],
    metadata: list[MetadataRecord],
) -> None:
    """Save merged source data and generated metadata to an XLSX workbook."""

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "merged_dataset"
    sheet.append(["proverb", "meaning", "category", "keywords", "english_translation"])

    for pair, record in zip(pairs, metadata):
        sheet.append(
            [
                pair.proverb,
                pair.meaning,
                record.category,
                ", ".join(record.keywords),
                record.english_translation,
            ]
        )

    workbook.save(output_path)


def generate_metadata_dataset(
    proverbs_path: Path,
    meanings_path: Path,
    output_path: Path,
    batch_size: int,
) -> None:
    """Run the full DOCX-to-XLSX metadata generation workflow."""

    validate_gemini_config()

    print("Reading Proverbs.docx...")
    proverbs = read_docx_paragraphs(proverbs_path)

    print("Reading Meanings.docx...")
    meanings = read_docx_paragraphs(meanings_path)

    print("Merging dataset...")
    pairs = merge_dataset(proverbs, meanings)

    print("Generating metadata...")
    generated_metadata: list[MetadataRecord] = []
    total = len(pairs)
    processed = 0

    for batch in tqdm(chunked(pairs, batch_size), total=(total + batch_size - 1) // batch_size):
        batch_metadata = generate_batch_metadata(batch)
        generated_metadata.extend(batch_metadata)
        for _ in batch_metadata:
            processed += 1
            print(f"{processed}/{total}")

    print("Saving merged_dataset.xlsx...")
    save_workbook(output_path, pairs, generated_metadata)
    print("Completed.")


def parse_args() -> argparse.Namespace:
    """Parse command-line arguments for the metadata generator."""

    parser = argparse.ArgumentParser(
        description="Generate Gemini metadata for Myanmar proverb DOCX datasets."
    )
    parser.add_argument(
        "--proverbs",
        type=Path,
        default=BACKEND_DIR / "Proverbs.docx",
        help="Path to Proverbs.docx.",
    )
    parser.add_argument(
        "--meanings",
        type=Path,
        default=BACKEND_DIR / "Meanings.docx",
        help="Path to Meanings.docx.",
    )
    parser.add_argument(
        "--output",
        type=Path,
        default=BACKEND_DIR / "merged_dataset.xlsx",
        help="Output XLSX path.",
    )
    parser.add_argument(
        "--batch-size",
        type=int,
        default=DEFAULT_BATCH_SIZE,
        help="Number of proverb-meaning pairs per Gemini request.",
    )
    return parser.parse_args()


def main() -> None:
    """CLI entry point."""

    args = parse_args()
    if args.batch_size <= 0:
        raise ValueError("--batch-size must be greater than zero.")

    configure_logging(LOG_FILE)
    generate_metadata_dataset(
        proverbs_path=args.proverbs,
        meanings_path=args.meanings,
        output_path=args.output,
        batch_size=args.batch_size,
    )


if __name__ == "__main__":
    main()
